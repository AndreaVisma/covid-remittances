"""Prepare IPUMS CPS occupational inputs for the remittance ABM.

Run: python oct-26/data-prep/cps-earnings-employment.py
Dependencies: numpy, pandas, openpyxl. No source records are edited.
See cps-earnings-employment-notes.md for definitions and interpretation.
"""
from pathlib import Path
import hashlib
import json
import re
import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# %% 1. Editable settings. Occupations describe jobs, not industrial sectors.
ROOT = Path('C:/git-projects/covid-remittances')
HERE = ROOT / 'oct-26/data-prep'
DATA = ROOT / 'data/cps_00001.dat'
CODEBOOK = HERE / 'cps_00001_codebook.txt'
OUT = HERE / 'outputs'
EXCEL = OUT / 'cps_earnings_employment.xlsx'
START_YEAR, END_YEAR = 2010, 2026
BASELINE_YEAR = 2019
MIN_AGE = 18                 # Match the ABM's adult sending eligibility.
CHUNK_ROWS = 250_000         # Read 3 GB in bounded chunks; retain Latino adults.
WEEKS_PER_MONTH = 52 / 12
MIN_EARN_N, MIN_EARN_EFFECTIVE_N = 30, 20
MIN_LABOUR_FORCE_N, MIN_UNEMPLOYED_N = 100, 5
MIN_STANDARD_CELL_N = 3
MIN_STANDARD_COVERAGE = 0.90
SAVE_SUMMARY_CSV = True
HISPANIC_CODES = [100, 102, 103, 104, 108, 109, 200, 300, 400, 500, 600, 610, 611, 612]
MEXICAN_CODES = [100, 102, 103, 104, 108, 109]
EMPLOYED_CODES = [10, 12]
UNEMPLOYED_CODES = [20, 21, 22]
NILF_CODES = [30, 31, 32, 33, 34, 35, 36]
WAGE_SALARY_CODES = [20, 21, 22, 23, 24, 25, 27, 28]
# Inclusive OCC2010 bounds. See exported code-level crosswalk for every label.
# Installation/repair is grouped with production; food/cleaning/protective/
# personal-care services are grouped with sales. Office support joins managers
# and professionals. Farming/fishing/forestry occupations define Agriculture.
OCCUPATION_RANGES = {
    'Agriculture': [(6000, 6130)],
    'Production/construction/transport': [(6200, 8965), (9000, 9750)],
    'Sales/personal services': [(3600, 4650), (4700, 4965)],
    'Professional/administrative': [(10, 3540), (5000, 5940)],
}
JOB_NAMES = list(OCCUPATION_RANGES) + ['Unknown']
POPULATIONS = ['Latino', 'Mexican origin']  # Overlapping; never add these rows.
AGE_BINS = [(18, 34), (35, 54), (55, 99)]
EDUC_BINS = {'High school or less': (2, 73), 'Some college or more': (80, 125)}


# %% 2. Read positions and decimal scales from the supplied extract codebook.
def read_layout():
    text = CODEBOOK.read_text(encoding='utf-8-sig')
    blocks = re.split(r'Variable: "([A-Z0-9_]+)"', text)
    layout, labels = {}, {}
    for name, block in zip(blocks[1::2], blocks[2::2]):
        start = int(re.search(r'Start Position:\s*(\d+)', block)[1]) - 1
        width = int(re.search(r'Width:\s*(\d+)', block)[1])
        decimals = int(re.search(r'Implied Decimal Places:\s*(\d+)', block)[1])
        layout[name] = (start, width, decimals)
        if name == 'OCC2010':
            labels = {int(k): v.strip() for k, v in re.findall(r'^\s*(\d+)\t([^\r\n]+)', block, re.M)}
    required = ['YEAR', 'MONTH', 'ASECFLAG', 'MISH', 'AGE', 'SEX', 'HISPAN', 'EMPSTAT',
                'OCC2010', 'CLASSWKR', 'EDUC', 'EARNWEEK2', 'EARNWT', 'WTFINL', 'CPSIDV']
    assert all(x in layout for x in required), 'Missing required variables in this extract.'
    return layout, labels


def occupation_codes(values):
    result = np.full(len(values), len(JOB_NAMES) - 1, dtype=np.int8)
    used = np.zeros(len(values), dtype=bool)
    for i, ranges in enumerate(OCCUPATION_RANGES.values()):
        mask = np.zeros(len(values), dtype=bool)
        for lo, hi in ranges:
            mask |= (values >= lo) & (values <= hi)
        assert not (mask & used).any(), 'Overlapping occupation definitions'
        result[mask] = i
        used |= mask
    return result


def load_people(layout):
    """Use the extract's fixed byte layout, not guessed whitespace columns.

    Filtering the ethnicity/age/date columns before parsing the rest keeps this
    faster and smaller than loading all 23 million records into a DataFrame.
    Blank ASECFLAG means Basic Monthly in this extract. Integer IDs never pass
    through floating point. Metadata decimal scaling is applied exactly once.
    """
    width = max(start + size for start, size, _ in layout.values())
    with DATA.open('rb') as f:
        first = f.readline()
    stride = len(first)
    assert len(first.rstrip(b'\r\n')) == width
    assert DATA.stat().st_size % stride == 0
    parts, raw_counts = [], []
    digest, seen = hashlib.sha256(), 0
    with DATA.open('rb') as f:
        while raw := f.read(stride * CHUNK_ROWS):
            digest.update(raw)
            n = len(raw) // stride
            assert len(raw) == n * stride
            assert np.all(np.ndarray((n,), dtype='S1', buffer=raw, offset=stride - 1, strides=(stride,)) == b'\n')
            def column(name, mask=None):
                start, size, decimals = layout[name]
                a = np.ndarray((n,), dtype=f'S{size}', buffer=raw, offset=start, strides=(stride,))
                if mask is not None:
                    a = a[mask]
                a = np.char.strip(a)
                a = np.where(a == b'', b'0', a).astype(np.int64)
                return a / 10 ** decimals if decimals else a
            year, month = column('YEAR'), column('MONTH')
            key, counts = np.unique(year * 100 + month, return_counts=True)
            raw_counts.append(pd.Series(counts, index=key))
            age, hispan, asec = column('AGE'), column('HISPAN'), column('ASECFLAG')
            keep = ((year >= START_YEAR) & (year <= END_YEAR) & (age >= MIN_AGE)
                    & (asec != 1) & np.isin(hispan, HISPANIC_CODES))
            names = ['YEAR', 'MONTH', 'MISH', 'AGE', 'SEX', 'HISPAN', 'EMPSTAT', 'OCC2010',
                     'CLASSWKR', 'EDUC', 'EARNWEEK2', 'EARNWT', 'WTFINL', 'CPSIDV']
            frame = pd.DataFrame({name: column(name, keep) for name in names})
            for name in names:
                if name not in ['EARNWEEK2', 'EARNWT', 'WTFINL', 'CPSIDV']:
                    frame[name] = frame[name].astype('int32')
            parts.append(frame)
            seen += n
            if len(parts) % 10 == 0:
                print(f'Read {seen:,} source rows...', flush=True)
    p = pd.concat(parts, ignore_index=True)
    raw_counts = pd.concat(raw_counts).groupby(level=0).sum().rename('source_records').rename_axis('ym').reset_index()
    assert p.MONTH.between(1, 12).all()
    p['ym'] = p.YEAR * 100 + p.MONTH
    p['mexican'] = p.HISPAN.isin(MEXICAN_CODES)
    p['job_code'] = occupation_codes(p.OCC2010.to_numpy())
    p['job_group'] = pd.Categorical.from_codes(p.job_code, JOB_NAMES)
    p['employed'] = p.EMPSTAT.isin(EMPLOYED_CODES)
    p['unemployed'] = p.EMPSTAT.isin(UNEMPLOYED_CODES)
    p['nilf'] = p.EMPSTAT.isin(NILF_CODES)
    p['civilian'] = p.employed | p.unemployed | p.nilf
    p['earn_eligible'] = (p.employed & p.CLASSWKR.isin(WAGE_SALARY_CODES)
                          & p.MISH.isin([4, 8]) & p.EARNWT.gt(0))
    p['earn_valid'] = p.earn_eligible & p.EARNWEEK2.ge(0) & p.EARNWEEK2.lt(999999.99)
    p['age_cell'] = np.select([p.AGE.between(lo, hi) for lo, hi in AGE_BINS], range(len(AGE_BINS)), default=-1)
    p['educ_cell'] = np.select([p.EDUC.between(lo, hi) for lo, hi in EDUC_BINS.values()], range(len(EDUC_BINS)), default=-1)
    p['std_cell'] = np.where((p.age_cell >= 0) & (p.educ_cell >= 0) & p.SEX.isin([1, 2]),
                             100 * p.SEX + 10 * p.age_cell + p.educ_cell, -1)
    # Later dynamic topcodes may exceed older plausible upper bounds. Retain
    # documented in-universe values rather than inventing an earnings cutoff.
    assert np.isfinite(p.loc[p.earn_valid, 'EARNWEEK2']).all()
    print('Maximum valid weekly earnings by year:',
          p.loc[p.earn_valid].groupby('YEAR').EARNWEEK2.max().to_dict(), flush=True)
    print(f'Retained {len(p):,} Latino adult observations.', flush=True)
    return p, raw_counts, digest.hexdigest()


# %% 3. Weighted statistics. Counts are observations, not distinct individuals.
def weighted_quantile(values, weights, qs):
    values, weights = np.asarray(values), np.asarray(weights)
    order = np.argsort(values)
    cumulative = np.cumsum(weights[order])
    return values[order][np.searchsorted(cumulative, np.asarray(qs) * cumulative[-1], side='left')]


def earnings_stats(g):
    if not len(g):
        return dict(earn_n=0, earn_effective_n=0., earn_weight=0., weekly_mean_usd=np.nan,
                    monthly_mean_usd=np.nan, monthly_p25_usd=np.nan, monthly_median_usd=np.nan,
                    monthly_p75_usd=np.nan, earnings_reliable=False)
    w, x = g.EARNWT.to_numpy(), g.EARNWEEK2.to_numpy()
    q = weighted_quantile(x, w, [.25, .5, .75]) * WEEKS_PER_MONTH
    mean, eff = np.average(x, weights=w), w.sum() ** 2 / (w @ w)
    return dict(earn_n=len(g), earn_effective_n=eff, earn_weight=w.sum(), weekly_mean_usd=mean,
                monthly_mean_usd=mean * WEEKS_PER_MONTH, monthly_p25_usd=q[0],
                monthly_median_usd=q[1], monthly_p75_usd=q[2],
                earnings_reliable=len(g) >= MIN_EARN_N and eff >= MIN_EARN_EFFECTIVE_N)


def employment_stats(g):
    g = g.loc[g.civilian & g.WTFINL.gt(0)]
    e, u = g.loc[g.employed], g.loc[g.unemployed]
    ew, uw = e.WTFINL.sum(), u.WTFINL.sum()
    lf = ew + uw
    return dict(adult_n=len(g), employed_n=len(e), unemployed_n=len(u),
        employed_weight=ew, unemployed_weight=uw, labour_force_weight=lf,
        adult_weight=g.WTFINL.sum(), unemployment_rate=uw / lf if lf else np.nan,
        unemployment_reliable=len(e) + len(u) >= MIN_LABOUR_FORCE_N and len(u) >= MIN_UNEMPLOYED_N)


def standard_references(p):
    b = p.loc[p.YEAR.eq(BASELINE_YEAR) & p.earn_valid & p.std_cell.ge(0)]
    references = {}
    for pop in POPULATIONS:
        s = b if pop == 'Latino' else b.loc[b.mexican]
        for job in JOB_NAMES + ['All occupations']:
            g = s if job == 'All occupations' else s.loc[s.job_group.eq(job)]
            w = g.groupby('std_cell').EARNWT.sum()
            references[pop, job] = w / w.sum()
    return references


def standardized_earnings(g, reference):
    # Direct standardization to baseline joint age/sex/education shares.
    # Missing small cells are NOT assigned zero earnings. Report supported mass.
    cells = g.loc[g.std_cell.ge(0)].groupby('std_cell')
    means = {k: np.average(c.EARNWEEK2, weights=c.EARNWT)
             for k, c in cells if len(c) >= MIN_STANDARD_CELL_N}
    common = reference.index.intersection(list(means))
    coverage = reference.loc[common].sum()
    value = (sum(reference[k] * means[k] for k in common) / coverage * WEEKS_PER_MONTH
             if coverage >= MIN_STANDARD_COVERAGE else np.nan)
    return value, coverage


def make_summaries(p):
    references = standard_references(p)
    earnings, employment, baseline = [], [], []
    for pop in POPULATIONS:
        population = p if pop == 'Latino' else p.loc[p.mexican]
        for ym, month in population.groupby('ym', sort=True):
            for job in JOB_NAMES + ['All occupations']:
                g = month if job == 'All occupations' else month.loc[month.job_group.eq(job)]
                eg = g.loc[g.earn_valid]
                row = dict(population=pop, year=int(ym // 100), month=int(ym % 100), job_group=job)
                es = earnings_stats(eg)
                std, coverage = standardized_earnings(eg, references[pop, job])
                earnings.append(dict(**row, **es, monthly_standardized_usd=std, standard_coverage=coverage,
                    eligible_earn_n=int(g.earn_eligible.sum()), missing_earn_n=int((g.earn_eligible & ~g.earn_valid).sum()),
                    fixed_topcode_share=(np.average(eg.EARNWEEK2.ge(2884.61), weights=eg.EARNWT)
                                         if len(eg) and ym <= 202303 else np.nan)))
                em = employment_stats(g)
                employment.append(dict(**row, **em,
                    employment_population_rate=(em['employed_weight'] / em['adult_weight']
                                                if job == 'All occupations' and em['adult_weight'] else np.nan),
                    participation_rate=(em['labour_force_weight'] / em['adult_weight']
                                        if job == 'All occupations' and em['adult_weight'] else np.nan)))
        b = population.loc[population.YEAR.eq(BASELINE_YEAR)]
        assert b.MONTH.nunique() == 12, 'Baseline requires all 12 months.'
        for job in JOB_NAMES + ['All occupations']:
            g = b if job == 'All occupations' else b.loc[b.job_group.eq(job)]
            es, em = earnings_stats(g.loc[g.earn_valid]), employment_stats(g)
            # Divide pooled weights by 12 for average monthly population totals.
            for stats in [es, em]:
                for key in stats:
                    if key.endswith('_weight'):
                        stats[key] /= 12
            baseline.append(dict(population=pop, year=BASELINE_YEAR, job_group=job, **es, **em))
        print(f'Summarized {pop}.', flush=True)
    return pd.DataFrame(baseline), pd.DataFrame(earnings), pd.DataFrame(employment)


# %% 4. Optional observed prior-earnings risk: 2019 earners linked to 2020.
def linked_earnings_risk(p):
    """Descriptive matched sample, NOT a representative unemployment input.

    MISH4 earners can be seen in MISH8 twelve months later. Weight by origin
    EARNWT only; no longitudinal/attrition-adjusted weights are in the extract.
    Never infer a former wage from a current missing/zero earnings field.
    """
    b = p.loc[p.YEAR.eq(BASELINE_YEAR) & p.earn_valid & p.MISH.eq(4) & p.CPSIDV.gt(0)].copy()
    later = p.loc[p.YEAR.eq(BASELINE_YEAR + 1) & p.MISH.eq(8) & p.CPSIDV.gt(0),
                  ['CPSIDV', 'MONTH', 'employed', 'unemployed', 'nilf']]
    assert not later.duplicated(['CPSIDV', 'MONTH']).any()
    matched = b.merge(later, on=['CPSIDV', 'MONTH'], how='left', suffixes=('', '_later'), validate='one_to_one', indicator=True)
    rows, thresholds = [], []
    for pop in POPULATIONS:
        all_base = p.loc[p.YEAR.eq(BASELINE_YEAR) & p.earn_valid]
        s = matched
        if pop == 'Mexican origin':
            all_base, s = all_base.loc[all_base.mexican], s.loc[s.mexican]
        cuts = weighted_quantile(all_base.EARNWEEK2, all_base.EARNWT, [.25, .5, .75])
        thresholds.append(dict(population=pop, year=BASELINE_YEAR,
            monthly_q25_usd=cuts[0] * WEEKS_PER_MONTH, monthly_q50_usd=cuts[1] * WEEKS_PER_MONTH,
            monthly_q75_usd=cuts[2] * WEEKS_PER_MONTH))
        s = s.copy()
        s['earnings_quartile_2019'] = np.searchsorted(cuts, s.EARNWEEK2, side='left') + 1
        for (month, job, quartile), g in s.groupby(['MONTH', 'job_group', 'earnings_quartile_2019'], observed=True):
            m = g.loc[g._merge.eq('both')]
            e = m.employed_later.eq(True)
            u = m.unemployed_later.eq(True)
            n = m.nilf_later.eq(True)
            lf_w = m.loc[e | u, 'EARNWT'].sum()
            matched_w = m.loc[e | u | n, 'EARNWT'].sum()
            reliable = len(m) >= MIN_EARN_N and int(u.sum()) >= MIN_UNEMPLOYED_N
            rows.append(dict(population=pop, year=BASELINE_YEAR + 1, month=int(month), baseline_job=str(job),
                earnings_quartile_2019=int(quartile), eligible_origin_n=len(g), matched_n=len(m),
                matched_share=len(m) / len(g), unemployed_n=int(u.sum()),
                matched_unemployment_rate=m.loc[u, 'EARNWT'].sum() / lf_w if lf_w else np.nan,
                matched_nonemployment_rate=m.loc[u | n, 'EARNWT'].sum() / matched_w if matched_w else np.nan,
                sufficient_sample=reliable, use_for_abm=False))
    return pd.DataFrame(rows), pd.DataFrame(thresholds)


# %% 5. Workbook documentation and readable, machine-friendly Excel export.
def method_rows():
    return [
        ('Source', 'IPUMS CPS extract cps_00001.dat; supplied extract codebook dated 2026-10-13. Original data remain unchanged.'),
        ('Source URL', 'https://cps.ipums.org/ ; https://cps.ipums.org/cps/citation.shtml'),
        ('Population', f'Hispanic/Latino civilians age {MIN_AGE}+ of any race. Mexican origin is an overlapping subset, not Mexican-born status. Do not add populations.'),
        ('Scope', 'Basic Monthly only (ASECFLAG != 1). Complete available history within configured year bounds. October 2025 is absent in this extract; do not fill missing months with zero.'),
        ('Baseline', 'Pooled 2019 observations weighted by EARNWT for earnings and WTFINL for labour force. Weight totals divided by 12 are average monthly totals. Counts are person-month observations.'),
        ('Units', 'Nominal USD. Monthly equivalents = usual gross weekly earnings * 52/12. These are neither net disposable budgets nor actual calendar-month earnings. No inflation adjustment.'),
        ('Earnings universe', 'Employed adult wage/salary workers in MISH 4/8, positive EARNWT, valid EARNWEEK2. Includes part-time and temporarily absent workers. Excludes self-employed and unpaid family workers.'),
        ('Missing and zero', '999999.99 is EARNWEEK2 NIU, not dollars. Valid zero earnings are retained. Nonworkers are excluded from conditional earnings, not assigned missing-code dollars.'),
        ('Weights / decimals', 'Use codebook field widths and implied decimals: EARNWEEK2 /100; EARNWT and WTFINL /10000. Applied once to raw fixed-width integers.'),
        ('Earnings documentation', 'https://cps.ipums.org/cps-action/variables/EARNWEEK2 ; https://cps.ipums.org/cps/outgoing_rotation_notes.shtml'),
        ('Occupation', 'OCC2010 harmonized codes, not IND industry. Crosswalk shows all code labels and assignments. Unknown jobs retained separately. Working student/pensioner cannot be identified from this extract.'),
        ('Unemployment', 'Weighted U/(E+U), with current occupation for employed and last reported occupation for unemployed. New entrants/unassignable jobs remain Unknown. Employment counts include self-employment.'),
        ('Nonparticipation', 'Overall employment/population and participation use all adult civilians. Occupation-specific NILF denominators are not used because past occupation coverage is incomplete.'),
        ('Employment documentation', 'https://cps.ipums.org/cps-action/variables/EMPSTAT ; https://cps.ipums.org/cps-action/variables/OCC ; https://cps.ipums.org/cps-action/variables/WTFINL'),
        ('Composition adjustment', 'Direct standardization within each population/job to 2019 joint age (18-34/35-54/55+), sex and education (<=high school / some college+) earnings-weight shares. Does not remove unobserved selection or finer occupational shifts.'),
        ('Standardization gaps', f'Cell mean requires n>={MIN_STANDARD_CELL_N}; supported baseline mass must be >= {MIN_STANDARD_COVERAGE:.0%}. Otherwise adjusted earnings blank. Supported shares renormalized; coverage exported. Unknown demographic cells excluded from reference.'),
        ('Earnings quality', f'Reliable flag requires raw n>={MIN_EARN_N} and Kish effective n>={MIN_EARN_EFFECTIVE_N}. Effective n reflects unequal weights only, not full survey-design variance.'),
        ('Unemployment quality', f'Reliable flag requires labour-force n>={MIN_LABOUR_FORCE_N} and unemployed n>={MIN_UNEMPLOYED_N}. Raw rates retained even when flagged. No design-based standard errors or confidence intervals.'),
        ('Unknown unemployment', 'Unknown often consists of unassignable unemployed/new entrants and nonparticipants, with no employed denominator. Its rate can be 100% by construction. Diagnostic only: never use it as a migrant job-loss probability.'),
        ('Topcoding', 'Pre-April 2023 fixed topcode share measures earnings >=2884.61. Later values blank because dynamic/phased topcoding is not reconstructed. Public-use topcodes retained; means are not topcode-corrected.'),
        ('Pandemic measurement', 'Official EMPSTAT used without speculative reclassification. Early-2020 temporary-absence misclassification can understate unemployment. https://www.bls.gov/blog/2020/update-on-the-misclassification-that-affected-the-unemployment-rate.htm'),
        ('Prior earnings risk', '2019 MISH4 earners linked by CPSIDV and calendar month to MISH8 in 2020. Quartiles from pooled 2019 earnings in each population. Boundary ties go to lower quartile; categories need not contain exactly 25%.'),
        ('Linked risk limitation', 'Origin EARNWT only, no longitudinal/attrition weights. Matched-sample descriptive rates, not population unemployment estimates. Coverage and sample flags exported; use_for_abm is False.'),
        ('Linking documentation', 'https://cps.ipums.org/cps-action/variables/CPSIDV ; https://cps.ipums.org/cps/cps_linking_documentation.shtml'),
        ('ABM use', 'Read baseline_2019 for static inputs and monthly_earnings/monthly_employment for time paths. Join by population/job_group/year/month. Check quality flags; never silently replace missing with zero.'),
        ('Capacity use', 'Earnings are gross upper bounds, not remittable disposable income. Unemployment rates are stocks, not job-loss probabilities. Do not simulate unemployment and also multiply employed earnings by 1-u.'),
        ('Refresh', 'Edit Python dictionaries/settings, then rerun script. Workbook cells are materialized analysis outputs for reliable pandas reading; Excel edits do not rerun microdata analysis.'),
    ]


def write_excel(tables):
    wb = Workbook()
    wb.remove(wb.active)
    for name, frame in tables.items():
        ws = wb.create_sheet(name)
        ws.append(list(frame.columns))
        for row in frame.itertuples(index=False, name=None):
            ws.append([None if pd.isna(x) else x.item() if isinstance(x, np.generic) else x for x in row])
        ws.freeze_panes = 'D2' if 'job_group' in frame else 'B2'
        ws.auto_filter.ref = ws.dimensions
        ws.sheet_view.showGridLines = False
        ws.row_dimensions[1].height = 44
        for cell in ws[1]:
            cell.font = Font(name='Arial', size=10, bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='243746')
            cell.alignment = Alignment(wrap_text=True, vertical='center')
        for j, col in enumerate(frame.columns, 1):
            width = 23
            if col in ['job_group', 'baseline_job']: width = 42
            if col in ['description', 'detail']: width = 105
            if col in ['label', 'topic']: width = 42
            if col == 'column': width = 38
            if col in ['matched_unemployment_rate', 'matched_nonemployment_rate']: width = 34
            if col == 'population': width = 20
            ws.column_dimensions[get_column_letter(j)].width = width
            for row in ws.iter_rows(min_row=2, min_col=j, max_col=j):
                cell = row[0]
                cell.font = Font(name='Arial', size=10)
                cell.alignment = Alignment(vertical='center', wrap_text=col in ['detail', 'label', 'description'])
                if col.endswith('_usd'): cell.number_format = '#,##0.00'
                elif col.endswith(('_rate', '_share', '_coverage')): cell.number_format = '0.0%'
                elif col.endswith(('_n', '_weight')): cell.number_format = '#,##0.0' if 'effective' in col else '#,##0'
        for i in range(2, ws.max_row + 1):
            ws.row_dimensions[i].height = 58 if name == 'methodology' else 30 if name in ['occupation_mapping', 'column_dictionary'] else 20
    wb['baseline_2019'].sheet_properties.tabColor = '243746'
    wb['methodology'].sheet_properties.tabColor = '80939E'
    wb.save(EXCEL)


def checks(p, baseline, earnings, employment, workbook_tables):
    # Independent accounting: disjoint occupation groups must sum to overall.
    keys = ['population', 'year', 'month']
    summed = employment.loc[employment.job_group.ne('All occupations')].groupby(keys)[['employed_weight', 'unemployed_weight']].sum()
    overall = employment.loc[employment.job_group.eq('All occupations')].set_index(keys)
    np.testing.assert_allclose(summed, overall.loc[summed.index, summed.columns], rtol=1e-10)
    np.testing.assert_allclose(earnings.monthly_mean_usd, earnings.weekly_mean_usd * WEEKS_PER_MONTH, equal_nan=True)
    assert employment.unemployment_rate.dropna().between(0, 1).all()
    assert not earnings.duplicated(keys + ['job_group']).any()
    assert not employment.duplicated(keys + ['job_group']).any()
    # Independently recompute one consequential baseline directly from persons.
    g = p.loc[p.YEAR.eq(BASELINE_YEAR) & p.earn_valid]
    expected = (g.EARNWT * g.EARNWEEK2).sum() / g.EARNWT.sum() * WEEKS_PER_MONTH
    got = baseline.loc[baseline.population.eq('Latino') & baseline.job_group.eq('All occupations'), 'monthly_mean_usd'].iloc[0]
    np.testing.assert_allclose(got, expected)
    book = load_workbook(EXCEL, read_only=True, data_only=True)
    for name, frame in workbook_tables.items():
        assert book[name].max_row == len(frame) + 1
        assert list(next(book[name].values)) == list(frame.columns)
    book.close()
    exported = pd.read_excel(EXCEL, sheet_name='baseline_2019')
    np.testing.assert_allclose(exported.monthly_mean_usd, baseline.monthly_mean_usd, equal_nan=True)
    print('PASS: subgroup accounting, conversion, baseline recomputation, keys, ranges and Excel read-back.', flush=True)


# %% 6. Main pipeline and audit files. No ABM file is read or changed.
def run():
    OUT.mkdir(parents=True, exist_ok=True)
    layout, labels = read_layout()
    people, coverage, sha = load_people(layout)
    baseline, earnings, employment = make_summaries(people)
    baseline['use_for_abm'] = (baseline.job_group.isin(OCCUPATION_RANGES)
                               & baseline.earnings_reliable & baseline.unemployment_reliable)
    risk, cuts = linked_earnings_risk(people)
    codes = sorted(set(labels) | set(people.OCC2010.unique()))
    assigned = occupation_codes(np.array(codes))
    mapping = pd.DataFrame(dict(occ2010=codes, label=[labels.get(c, 'No label in extract codebook') for c in codes],
                                job_group=[JOB_NAMES[i] for i in assigned]))
    coverage = coverage.merge(people.groupby('ym').size().rename('latino_adult_records'), on='ym', how='left')
    coverage['year'], coverage['month'] = coverage.ym // 100, coverage.ym % 100
    coverage['within_analysis_years'] = coverage.year.between(START_YEAR, END_YEAR)
    expected = pd.period_range(f'{START_YEAR}-01', f'{END_YEAR}-12', freq='M')
    last = pd.Period(f'{int(coverage.year.max())}-{int(coverage.loc[coverage.year.eq(coverage.year.max()), "month"].max())}', freq='M')
    observed = set(coverage.ym)
    missing = [str(d) for d in expected if d <= last and d.year * 100 + d.month not in observed]
    method = pd.DataFrame(method_rows(), columns=['topic', 'detail'])
    method.loc[len(method)] = ['Missing months in requested observed span', ', '.join(missing) or 'None']
    method.loc[len(method)] = ['Composition reference missing demographics', str(int((people.YEAR.eq(BASELINE_YEAR) & people.earn_valid & people.std_cell.lt(0)).sum()))]
    # Column dictionary makes the Excel interface readable without inspecting code.
    tables = dict(baseline_2019=baseline, monthly_earnings=earnings, monthly_employment=employment,
                  linked_earnings_risk=risk, earnings_cutoffs=cuts, occupation_mapping=mapping,
                  sample_coverage=coverage, methodology=method)
    descriptions = {
        'earn_n': 'Unweighted number of valid earnings observations.',
        'earn_effective_n': 'Kish effective sample size: sum(w)^2/sum(w^2), not design-adjusted.',
        'earn_weight': 'Sum of earnings weights (baseline divided by 12).',
        'monthly_mean_usd': 'Weighted usual weekly mean multiplied by 52/12, nominal gross USD.',
        'monthly_standardized_usd': 'Mean with 2019 demographic cell shares; blank if coverage insufficient.',
        'standard_coverage': 'Baseline cell weight share supported by this month, conditional on known baseline cells.',
        'unemployment_rate': 'WTFINL-weighted unemployed / (employed + unemployed).',
        'unemployment_reliable': 'Meets configured labour-force and unemployed observation thresholds.',
        'earnings_reliable': 'Meets configured raw and effective earnings sample thresholds.',
        'employment_population_rate': 'Overall employed / adult civilians; deliberately blank by occupation.',
        'participation_rate': 'Overall labour force / adult civilians; deliberately blank by occupation.',
        'fixed_topcode_share': 'Weighted share >=2884.61 through March 2023; later blank.',
        'matched_unemployment_rate': 'Origin-EARNWT weighted U/(E+U) in linked sample only.',
        'matched_nonemployment_rate': 'Origin-EARNWT weighted (U+NILF)/(E+U+NILF) in linked sample only.',
        'use_for_abm': 'Baseline: known job and both reliability flags pass. Linked risk: always False, no attrition-adjusted weights.',
    }
    tables['column_dictionary'] = pd.DataFrame(descriptions.items(), columns=['column', 'description'])
    tables['methodology'] = tables.pop('methodology')
    if SAVE_SUMMARY_CSV:
        for name, frame in tables.items():
            frame.to_csv(OUT / f'cps_{name}.csv', index=False)
    write_excel(tables)
    checks(people, baseline, earnings, employment, tables)
    report = dict(source=str(DATA), source_sha256=sha, source_records=int(coverage.source_records.sum()),
        latino_adult_records=len(people), baseline_year=BASELINE_YEAR, missing_months=missing,
        workbook=str(EXCEL), sheets={k: len(v) for k, v in tables.items()},
        occupation_ranges=OCCUPATION_RANGES, weeks_per_month=WEEKS_PER_MONTH,
        settings=dict(start_year=START_YEAR, end_year=END_YEAR, min_age=MIN_AGE,
            min_earn_n=MIN_EARN_N, min_earn_effective_n=MIN_EARN_EFFECTIVE_N,
            min_labour_force_n=MIN_LABOUR_FORCE_N, min_unemployed_n=MIN_UNEMPLOYED_N,
            min_standard_cell_n=MIN_STANDARD_CELL_N, min_standard_coverage=MIN_STANDARD_COVERAGE))
    (OUT / 'cps_run_settings.json').write_text(json.dumps(report, indent=2), encoding='utf-8')
    print(baseline[['population', 'job_group', 'monthly_mean_usd', 'unemployment_rate']].to_string(index=False), flush=True)
    print(f'Saved {EXCEL}', flush=True)
    return report


if __name__ == '__main__':
    run()
