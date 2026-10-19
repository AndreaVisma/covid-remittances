"""Synthetic input generation, strict CSV loading, and read-only target extraction."""
from pathlib import Path
import json
import numpy as np
import pandas as pd
import openpyxl


def generate_inputs(config, folder):
    """Never reads outcomes. Refuses to overwrite existing researcher inputs."""
    folder = Path(folder)
    if folder.exists() and any(folder.iterdir()):
        raise ValueError(f"Input directory is not empty: {folder}. Choose a new directory.")
    folder.mkdir(parents=True, exist_ok=True)
    rng = np.random.default_rng(config['synthetic_seed'])
    n = config['n_agents']
    if n < 2 or config['represented_population'] <= 0:
        raise ValueError('Require at least two agents and a positive represented population.')
    # Equal weights are essential: connected groups represent replicated families.
    agents = pd.DataFrame({
        'agent_id': np.arange(n), 'recipient_id': np.arange(n) // 3,
        'sector': rng.choice(['exposed', 'other'], n, p=[0.35, 0.65]),
        'reference_income_usd': rng.lognormal(np.log(2800), 0.55, n),
        'essential_cost_usd': rng.uniform(1000, 1700, n),
        'base_probability': rng.uniform(0.15, 0.4, n),
        'tie_strength': rng.uniform(0.4, 1.0, n),
        'initial_employed': (rng.random(n) < 0.92).astype(int),
        'initial_willing': (rng.random(n) < 0.275).astype(int),
        'weight': np.full(n, config['represented_population'] / n)
    })
    dates = pd.date_range(config['start'], config['end'], freq='MS')
    shocks, needs = [], []
    recipients = agents.recipient_id.unique()
    vulnerability = rng.uniform(0.6, 1.0, len(recipients))
    for date in dates:
        months = (date.year - 2020) * 12 + date.month - 4
        crisis = 0.0 if months < 0 else np.exp(-months / 14)
        for sector in ['exposed', 'other']:
            severity = 1.0 if sector == 'exposed' else 0.4
            shocks.append([date, sector, 0.01 + 0.22 * crisis * severity,
                           0.12, 1 - 0.12 * crisis * severity])
        needs.extend([date, int(r), float(crisis * v)] for r, v in zip(recipients, vulnerability))
    agents.to_csv(folder / 'agents.csv', index=False)
    pd.DataFrame(shocks, columns=['date', 'sector', 'job_loss_probability',
                                 'job_find_probability', 'wage_multiplier']).to_csv(folder / 'employment.csv', index=False)
    pd.DataFrame(needs, columns=['date', 'recipient_id', 'need']).to_csv(folder / 'needs.csv', index=False)
    (folder / 'provenance.json').write_text(json.dumps({
        'kind': 'synthetic', 'description': 'All inputs invented; independent of target outcomes.',
        'generator_seed': config['synthetic_seed'], 'config': config
    }, indent=2), encoding='utf-8')


def load_inputs(folder, config):
    folder = Path(folder)
    a = pd.read_csv(folder / 'agents.csv')
    e = pd.read_csv(folder / 'employment.csv', parse_dates=['date'])
    h = pd.read_csv(folder / 'needs.csv', parse_dates=['date'])
    required = {
        'agents': (a, ['agent_id', 'recipient_id', 'sector', 'reference_income_usd',
                      'essential_cost_usd', 'base_probability', 'tie_strength',
                      'initial_employed', 'initial_willing', 'weight']),
        'employment': (e, ['date', 'sector', 'job_loss_probability', 'job_find_probability', 'wage_multiplier']),
        'needs': (h, ['date', 'recipient_id', 'need'])}
    for name, (frame, cols) in required.items():
        missing = set(cols) - set(frame.columns)
        if missing:
            raise ValueError(f'{name}: missing columns {missing}')
        if frame[cols].isna().any().any() or frame.empty:
            raise ValueError(f'{name}: missing values or empty table')
        numeric = [c for c in cols if c not in ['agent_id', 'recipient_id', 'sector', 'date']]
        if not np.isfinite(frame[numeric].to_numpy(dtype=float)).all():
            raise ValueError(f'{name}: nonfinite numeric values')
    for frame, keys in [(a, ['agent_id']), (e, ['date', 'sector']), (h, ['date', 'recipient_id'])]:
        if frame.duplicated(keys).any():
            raise ValueError(f'Duplicate keys: {keys}')
    for frame, cols in [(a, ['base_probability', 'tie_strength']),
                        (e, ['job_loss_probability', 'job_find_probability']), (h, ['need'])]:
        if not frame[cols].apply(lambda s: s.between(0, 1)).all().all():
            raise ValueError(f'Values outside [0, 1]: {cols}')
    if not a[['initial_employed', 'initial_willing']].isin([0, 1]).all().all():
        raise ValueError('Initial states must be 0 or 1')
    if (a[['reference_income_usd', 'essential_cost_usd']] < 0).any().any() or (e.wage_multiplier < 0).any():
        raise ValueError('Negative monetary inputs or wage multipliers')
    if (a.weight <= 0).any() or not np.allclose(a.weight, a.weight.iloc[0]):
        raise ValueError('This network prototype requires positive equal weights; resample complete groups.')
    dates = pd.date_range(config['start'], config['end'], freq='MS')
    for frame, key, values in [(e, 'sector', a.sector.unique()), (h, 'recipient_id', a.recipient_id.unique())]:
        if not (frame.date.dt.day == 1).all():
            raise ValueError('Dates must be month starts')
        expected = pd.MultiIndex.from_product([dates, values], names=['date', key])
        if len(expected.difference(pd.MultiIndex.from_frame(frame[['date', key]]))):
            raise ValueError(f'Incomplete monthly coverage: {key}')
    a = a.sort_values('agent_id').reset_index(drop=True)
    return a, e, h


def load_target(path, config):
    path = Path(path)
    if not np.isfinite(config['target_multiplier']) or config['target_multiplier'] <= 0:
        raise ValueError('target_multiplier must be finite and positive')
    if not path.exists():
        raise FileNotFoundError(f'Target workbook not found: {path}. Pass --target with the correct path.')
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    audit = {'path': str(path.resolve()), 'sheets': {},
             'currency_assumption': config['target_currency_assumption'],
             'coverage_warning': 'No sending-country or migrant-population identification in workbook; aggregate Mexico comparison only.',
             'multiplier_to_usd_assumed': config['target_multiplier']}
    for sheet in workbook:
        audit['sheets'][sheet.title] = {'rows': sheet.max_row, 'columns': sheet.max_column}
    sheet = workbook[config['target_sheet']]
    rows = list(sheet.values)
    frame = pd.DataFrame(rows[1:], columns=rows[0]).dropna(how='all')
    workbook.close()
    audit['columns'] = [str(c) for c in frame.columns]
    audit['missing_by_column'] = {str(k): int(v) for k, v in frame.isna().sum().items()}
    dates = pd.to_datetime(frame['date'], errors='raise')
    values = pd.to_numeric(frame[config['target_value_column']], errors='raise')
    if dates.isna().any() or dates.duplicated().any() or not (dates.dt.day == 1).all():
        raise ValueError('Target has missing, duplicate, or non-month-start dates')
    if not np.isfinite(values).all() or (values < 0).any():
        raise ValueError('Target has invalid amounts')
    target = pd.DataFrame({'date': dates, 'observed_usd': values * config['target_multiplier']}).sort_values('date')
    gaps = pd.date_range(target.date.min(), target.date.max(), freq='MS').difference(target.date)
    audit.update(start=str(target.date.min().date()), end=str(target.date.max().date()),
                 observations=len(target), missing_months=[str(d.date()) for d in gaps])
    return target, audit
